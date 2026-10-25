import argparse
import base64
import io
import json
import mimetypes
import re
import secrets
import socket
import threading
import urllib.parse
import urllib.request
import webbrowser
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from .catalog import TEMPLATES
from .engine import Engine
from .store import Store
from .assistant import Assistant, editable_path
from .recipes import skill_catalog
from .workspace import Workspace, digest
from .intake import Intake
from .studio import Studio

APP_ROOT=Path(__file__).resolve().parent.parent

class Application:
    def __init__(self,data_root):
        self.store=Store(data_root)
        self.engine=Engine(self.store)
        self.assistant=Assistant(self.engine)
        self.workspace=Workspace(self.engine)
        self.intake=Intake(self)
        self.studio=Studio(self)
        self.token=secrets.token_urlsafe(32)

    def draft_only(self,pid):
        project=self.store.get(pid)
        if project['status']!='draft': raise ValueError('资料与文献需在启动流程前添加；新资料请用于新项目')
        return project

    def upload(self,pid,value):
        self.draft_only(pid)
        name=value.get('name','')
        if not isinstance(name,str) or not name or len(name)>160 or re.search(r'[\\/:*?"<>|\x00-\x1f]',name) or name.startswith('.'):
            raise ValueError('文件名无效')
        suffix=Path(name).suffix.lower()
        if suffix not in ('.pdf','.docx','.txt','.md','.csv','.json','.bib'): raise ValueError('支持 PDF、DOCX、TXT、MD、CSV、JSON、BIB')
        files=self.store.files(pid)
        if len([f for f in files if f['path'].startswith('materials/')])>=100: raise ValueError('每项目最多 100 个资料文件（含提取文本）')
        raw=base64.b64decode(value.get('content',''),validate=True)
        if len(raw)>12_000_000: raise ValueError('单个资料上限 12 MB')
        if sum(f['size'] for f in files)+len(raw)>150_000_000: raise ValueError('项目资料超过 150 MB')
        dest=self.store.path(pid,'materials/'+name)
        if dest.exists(): raise ValueError('同名资料已存在，请修改文件名')
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(raw))
            text='\n'.join((p.extract_text() or '')[:15000] for p in reader.pages[:80])[:200000]
        elif suffix=='.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>40_000_000: raise ValueError('Word 解压大小超限')
            doc=Document(io.BytesIO(raw))
            text='\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])[:200000]
        else:
            try: text=raw.decode('utf-8-sig')[:200000]
            except UnicodeDecodeError: text=raw.decode('gb18030',errors='replace')[:200000]
        dest.write_bytes(raw)
        if suffix in ('.pdf','.docx'):
            self.store.write(pid,'materials/'+name+'.extracted.txt',text or '未能提取文本，可能是扫描件；本版本不提供 OCR。')
        self.store.event(pid,'info','导入资料：'+name+'；文本供模型参考，文件可能截断，不等于来源已核验')
        return {'name':name,'characters':len(text)}

    def search(self,pid,value):
        self.draft_only(pid)
        query=str(value.get('query','')).strip()[:500]
        if not query: raise ValueError('请输入文献关键词')
        url='https://api.crossref.org/works?'+urllib.parse.urlencode({'query':query,'rows':8})
        req=urllib.request.Request(url,headers={'User-Agent':'ResearchDesk/0.1 (personal research metadata lookup)'})
        try:
            with urllib.request.urlopen(req,timeout=25) as r: data=json.loads(r.read(4_000_000))
            items=data['message']['items']
        except Exception: raise ValueError('Crossref 检索失败，请检查网络；你也可以上传已核验的文献资料') from None
        sources=[]
        for item in items:
            sources.append({'title':' '.join(item.get('title',[])),'doi':item.get('DOI',''),
                'authors':[(' '.join([a.get('given',''),a.get('family','')])).strip() for a in item.get('author',[])[:20]],
                'year':(item.get('published',{}).get('date-parts') or [[None]])[0][0],
                'url':item.get('URL',''),'source':'Crossref metadata; full text not retrieved'})
        self.store.write(pid,'references/sources.json',json.dumps({'query':query,'items':sources,'notice':'检索结果仅为候选元数据，需人工核验相关性与全文'},ensure_ascii=False,indent=2))
        def clean(v): return re.sub(r'[{}\\]','',str(v))
        bib='\n\n'.join('@article{ref'+str(i+1)+',\n  title={'+clean(x['title'])+'},\n  doi={'+clean(x['doi'])+'},\n  year={'+clean(x['year'] or '')+'}\n}' for i,x in enumerate(sources))
        self.store.write(pid,'references/references.bib',bib or '% 没有检索结果')
        self.store.event(pid,'info',f'Crossref 返回 {len(sources)} 条候选元数据，未获取全文')
        return sources

class Handler(BaseHTTPRequestHandler):
    server_version='ResearchDesk/0.3'
    def log_message(self,fmt,*args): pass
    @property
    def app(self): return self.server.app

    def headers_out(self,status,mime,length,disposition=None):
        self.send_response(status)
        self.send_header('Content-Type',mime)
        self.send_header('Content-Length',str(length))
        self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        self.send_header('X-Frame-Options','DENY')
        self.send_header('Referrer-Policy','no-referrer')
        self.send_header('Content-Security-Policy',"default-src 'none'; style-src 'unsafe-inline'; img-src data:; sandbox;" if mime=='image/svg+xml' else "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data: blob:; object-src 'none'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
        if disposition: self.send_header('Content-Disposition',disposition)
        self.end_headers()

    def respond(self,value,status=200):
        raw=json.dumps(value,ensure_ascii=False).encode('utf-8')
        self.headers_out(status,'application/json; charset=utf-8',len(raw))
        self.wfile.write(raw)

    def host_ok(self):
        return self.headers.get('Host') in (f'127.0.0.1:{self.server.server_port}',f'localhost:{self.server.server_port}')

    def do_GET(self):
        if not self.host_ok(): return self.respond({'error':'Host 不允许'},403)
        url=urllib.parse.urlsplit(self.path)
        path=url.path
        args=urllib.parse.parse_qs(url.query)
        try:
            if path=='/api/bootstrap':
                return self.respond({'token':self.app.token,'version':'0.3.0','templates':TEMPLATES,'projects':self.app.store.list(),'settings':self.app.store.settings(),'profiles':self.app.store.profiles(),'skills':skill_catalog(),'intakes':self.app.intake.list()})
            if path=='/api/intake':return self.respond(self.app.intake.get(args.get('id',[''])[0]))
            if path=='/api/skill':return self.respond(self.app.workspace.skill(args.get('name',[''])[0]))
            if path=='/api/intakes':return self.respond(self.app.intake.list())
            if path=='/api/versions':return self.respond(self.app.workspace.list(args.get('id',[''])[0]))
            if path=='/api/diagrams':return self.respond(self.app.studio.diagrams())
            if path=='/api/diagram':return self.respond(self.app.studio.diagram(args.get('id',[''])[0],args.get('name',[''])[0]))
            if path=='/api/surveys':return self.respond(self.app.studio.survey_list())
            if path=='/api/survey':return self.respond(self.app.studio.survey_get(args.get('id',[''])[0]))
            if path=='/api/survey.csv':
                raw=self.app.studio.survey_csv(args.get('id',[''])[0]);self.headers_out(200,'text/csv; charset=utf-8',len(raw),'attachment; filename=responses.csv');self.wfile.write(raw);return
            if path=='/api/projects': return self.respond(self.app.store.list())
            if path=='/api/project': return self.respond(self.app.engine.detail(args.get('id',[''])[0]))
            if path=='/api/file':
                f=self.app.store.path(args.get('id',[''])[0],args.get('path',[''])[0])
                if not f.is_file(): raise ValueError('文件不存在')
                if f.stat().st_size>2_000_000: raise ValueError('此文件过大，请直接下载')
                if f.suffix not in ('.md','.txt','.csv','.json','.py','.bib','.log','.svg','.html','.css','.js','.tex','.sql','.drawio'): raise ValueError('此文件请下载查看')
                return self.respond({'path':args['path'][0],'content':f.read_text('utf-8',errors='replace'),'sha256':digest(f.read_bytes())})
            if path=='/api/download':
                f=self.app.store.path(args.get('id',[''])[0],args.get('path',[''])[0])
                if not f.is_file(): raise ValueError('文件不存在')
                mime=mimetypes.guess_type(f.name)[0] or 'application/octet-stream'
                disposition="attachment; filename*=UTF-8''"+urllib.parse.quote(f.name)
                if args.get('inline')==['1'] and f.suffix in ('.png','.svg'): disposition=None
                self.headers_out(200,mime,f.stat().st_size,disposition)
                with f.open('rb') as source:
                    while chunk:=source.read(65536): self.wfile.write(chunk)
                return
            static={'/':'index.html','/app.js':'app.js','/style.css':'style.css','/favicon.svg':'favicon.svg','/workbench.js':'workbench.js','/workbench.css':'workbench.css','/v3.js':'v3.js','/v3.css':'v3.css'}
            if path not in static: return self.respond({'error':'未找到'},404)
            f=APP_ROOT/'static'/static[path]
            raw=f.read_bytes()
            self.headers_out(200,{'html':'text/html; charset=utf-8','js':'application/javascript; charset=utf-8','css':'text/css; charset=utf-8','svg':'image/svg+xml'}[f.suffix[1:]],len(raw))
            self.wfile.write(raw)
        except (ValueError,KeyError,OSError) as exc: self.respond({'error':str(exc)[:1000]},400)

    def do_POST(self):
        try:
            length=int(self.headers.get('Content-Length','0'))
            if not 0<length<=17_000_000: return self.respond({'error':'请求为空或超过大小限制'},413)
            self.connection.settimeout(30)
            request_body=self.rfile.read(length)
            if len(request_body)!=length: return self.respond({'error':'请求正文不完整'},400)
        except (ValueError,OSError): return self.respond({'error':'请求正文读取失败'},400)
        if not self.host_ok() or not secrets.compare_digest(self.headers.get('X-ResearchDesk-Token',''),self.app.token):
            return self.respond({'error':'请在本机工作台页面操作（令牌无效）'},403)
        origin=self.headers.get('Origin')
        if origin and origin not in (f'http://127.0.0.1:{self.server.server_port}',f'http://localhost:{self.server.server_port}'):
            return self.respond({'error':'来源不允许'},403)
        try:
            value=json.loads(request_body)
            if not isinstance(value,dict): raise ValueError('请求必须是 JSON 对象')
            path=urllib.parse.urlsplit(self.path).path
            pid=value.get('id','')
            if path=='/api/skill':return self.respond(self.app.workspace.save_skill(value.get('name'),value.get('content'),value.get('sha256')))
            if path=='/api/intake/new':return self.respond(self.app.intake.new())
            if path=='/api/intake/upload':return self.respond(self.app.intake.upload(pid,value))
            if path=='/api/intake/send':return self.respond(self.app.intake.send(pid,value.get('text',''),value.get('turn_id',''),value.get('template','auto'),bool(value.get('local',False))))
            if path=='/api/intake/edit':return self.respond(self.app.intake.edit(pid,int(value.get('revision',0)),value.get('proposal')))
            if path=='/api/intake/confirm':return self.respond(self.app.intake.confirm(pid,int(value.get('revision',0))))
            if path=='/api/intake/stop':return self.respond(self.app.intake.stop(pid))
            if path=='/api/project/settings':return self.respond(self.app.workspace.patch_settings(pid,value.get('patch'),value.get('revision',1)))
            if path=='/api/versions/save':return self.respond(self.app.workspace.capture(pid,value.get('label','手动保存')))
            if path=='/api/versions/preview':return self.respond(self.app.workspace.preview(pid,value.get('version_id')))
            if path=='/api/versions/restore':return self.respond(self.app.workspace.restore(pid,value.get('version_id'),value.get('token')))
            if path=='/api/diagram/save':return self.respond(self.app.studio.save_diagram(pid,value.get('name'),value.get('document'),value.get('token')))
            if path=='/api/survey/create':return self.respond(self.app.studio.survey_create(value))
            if path=='/api/survey/answer':return self.respond(self.app.studio.answer(pid,value.get('answers')))
            if path=='/api/survey/toggle':return self.respond(self.app.studio.survey_toggle(pid,int(value.get('revision',0))))
            if path=='/api/survey/import':return self.respond(self.app.studio.survey_import(value.get('survey_id'),pid))
            if path=='/api/projects':
                return self.respond(self.app.store.create(value.get('title',''),value.get('template','research'),value.get('goal',''),value.get('mode','api'),value.get('options')))
            if path=='/api/settings': return self.respond(self.app.store.save_settings(value))
            if path=='/api/profiles': return self.respond(self.app.store.save_profile(value))
            if path=='/api/chat': return self.respond(self.app.assistant.send(pid,value.get('text',''),value.get('mode','auto'),value.get('turn_id','')))
            if path=='/api/chat/stop': return self.respond(self.app.assistant.stop(pid))
            if path=='/api/task': return self.respond(self.app.assistant.operate(pid,value.get('turn_id'),value.get('action'),value.get('revision',1),value.get('proposal')))
            if path=='/api/rollback': return self.respond(self.app.assistant.rollback(pid,int(value.get('snapshot_id',0))))
            if path=='/api/rerun':
                self.app.engine.rerun(pid,int(value.get('index',-1)))
                return self.respond({'ok':True})
            if path=='/api/run':
                self.app.engine.start(pid,value.get('action','continue'))
                return self.respond({'ok':True})
            if path=='/api/pause':
                self.app.engine.pause(pid)
                return self.respond({'ok':True})
            if path=='/api/upload': return self.respond(self.app.upload(pid,value))
            if path=='/api/search': return self.respond(self.app.search(pid,value))
            if path=='/api/file':
                p=self.app.store.get(pid)
                detail=self.app.engine.detail(pid)
                if p['status']=='running' or detail['busy'] or detail.get('chat_busy'): raise ValueError('运行时不能编辑文件')
                relative=value.get('path','')
                f=self.app.store.path(pid,relative)
                if relative.startswith('code/') and p['status']!='awaiting_execution': raise ValueError('仅可在等待计算确认时修改代码')
                if not editable_path(relative) or not f.exists(): raise ValueError('此文件不支持编辑')
                content=value.get('content')
                if not isinstance(content,str) or not content.strip(): raise ValueError('内容不能为空')
                self.app.workspace.edit(pid,relative,content,value.get('sha256',digest(f.read_bytes())))
                self.app.store.event(pid,'info','用户编辑文件：'+relative+('；请重新导出以更新 Word/PDF' if p['status']=='completed' else ''))
                return self.respond({'ok':True})
            return self.respond({'error':'未找到'},404)
        except (ValueError,KeyError,TypeError,OSError) as exc: self.respond({'error':str(exc)[:1500]},400)
        except Exception:
            self.respond({'error':'操作失败，请检查文件格式或本机依赖；项目数据已保留'},500)

def main():
    parser=argparse.ArgumentParser(description='ResearchDesk personal research workbench')
    parser.add_argument('--port',type=int,default=18765)
    parser.add_argument('--data-dir',default=str(APP_ROOT/'data'))
    parser.add_argument('--no-browser',action='store_true')
    args=parser.parse_args()
    try: server=ThreadingHTTPServer(('127.0.0.1',args.port),Handler)
    except OSError:
        print(f'端口 {args.port} 已占用。请打开 http://127.0.0.1:{args.port} 或使用 --port 更换端口。',flush=True)
        return 1
    server.app=Application(args.data_dir)
    server.daemon_threads=True
    address=f'http://127.0.0.1:{server.server_port}'
    print(f'ResearchDesk 0.3.0 已启动：{address}\n数据：{Path(args.data_dir).resolve()}\n关闭此窗口将停止服务。',flush=True)
    if not args.no_browser: threading.Timer(.6,lambda:webbrowser.open(address)).start()
    try: server.serve_forever()
    except KeyboardInterrupt: pass
    finally: server.server_close()
    return 0

if __name__=='__main__': raise SystemExit(main())
