import json
import math
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile
import base64
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from researchdesk import demo, providers
from researchdesk.artifacts import validate_metrics
from researchdesk.catalog import BY_ID
from researchdesk.engine import Engine
from researchdesk.server import Application, Handler
from researchdesk.store import Store

class FakeAPI(BaseHTTPRequestHandler):
    records=[]
    def log_message(self,*args): pass
    def do_POST(self):
        body=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        self.records.append({'path':self.path,'body':body,'headers':dict(self.headers)})
        prompt=body.get('input') or body.get('messages',[{}])[-1].get('content','')
        try:
            request=json.loads(prompt)
            step=BY_ID[request['step']]
            files=[]
            for f in request.get('required_outputs',step.outputs):
                content=demo.CODE if f.endswith('.py') else '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><text x="10" y="30">Test</text></svg>' if f.endswith('.svg') else '# 本地假 API 测试\n\n模拟文本，非真实科研结论。'
                files.append({'path':f,'content':content})
            text=json.dumps({'summary':'fake API response','files':files},ensure_ascii=False)
        except (ValueError,KeyError): text='test response'
        if self.path.endswith('/chat/completions'):
            data={'choices':[{'message':{'content':text},'finish_reason':'stop'}],'usage':{'prompt_tokens':100,'completion_tokens':50}}
        elif self.path.endswith('/responses'):
            data={'status':'completed','output':[{'type':'reasoning','summary':[]},{'type':'message','content':[{'type':'output_text','text':text}]}],'usage':{'input_tokens':100,'output_tokens':50}}
        else:
            data={'content':[{'type':'text','text':text}],'stop_reason':'end_turn','usage':{'input_tokens':100,'output_tokens':50}}
        raw=json.dumps(data).encode()
        self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)

class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fake=ThreadingHTTPServer(('127.0.0.1',0),FakeAPI)
        threading.Thread(target=cls.fake.serve_forever,daemon=True).start()
        cls.base=f'http://127.0.0.1:{cls.fake.server_port}/v1'
    @classmethod
    def tearDownClass(cls): cls.fake.shutdown();cls.fake.server_close()
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.store=Store(self.tmp.name)
        self.engine=Engine(self.store)
    def tearDown(self): self.tmp.cleanup()
    def project(self,mode='demo',template='research'):
        return self.store.create('测试项目',template,'研究目标',mode)['id']
    def wait(self,pid,timeout=25):
        end=time.monotonic()+timeout
        while time.monotonic()<end:
            with self.engine.guard: busy=pid in self.engine.active
            if not busy: return self.store.get(pid)
            time.sleep(.025)
        self.fail('background task timeout')
    def run_to_end(self,pid):
        self.engine.start(pid)
        p=self.wait(pid)
        self.assertEqual(p['status'],'review',p)
        self.assertTrue(self.store.path(pid,'docs/plan.md').exists())
        self.engine.start(pid)
        p=self.wait(pid)
        self.assertEqual(p['status'],'awaiting_execution',p)
        self.assertFalse(self.store.path(pid,'results/metrics.json').exists())
        self.engine.start(pid,'execute')
        p=self.wait(pid)
        self.assertEqual(p['status'],'review',p)
        receipt=json.loads(self.store.path(pid,'results/execution.json').read_text('utf-8'))
        self.assertEqual(receipt['status'],'executed')
        self.assertEqual(len(receipt['code_sha256']),64)
        self.engine.start(pid)
        p=self.wait(pid)
        self.assertEqual(p['status'],'completed',p)
        return p

    def test_demo_end_to_end_real_execution_exports(self):
        pid=self.project()
        with patch('researchdesk.providers.call',side_effect=AssertionError('Demo must not call an API')):
            self.run_to_end(pid)
        metrics=validate_metrics(json.loads(self.store.path(pid,'results/metrics.json').read_text('utf-8')))
        self.assertLess(metrics['metrics'][0]['value'],metrics['metrics'][-1]['value'])
        from docx import Document
        from pypdf import PdfReader
        self.assertGreater(len(Document(self.store.path(pid,'exports/paper.docx')).paragraphs),15)
        pdf=PdfReader(self.store.path(pid,'exports/paper.pdf'))
        self.assertGreater(len(pdf.pages),0)
        self.assertIn('仿真',''.join(p.extract_text() for p in pdf.pages))
        with zipfile.ZipFile(self.store.path(pid,'exports/project.zip')) as z:
            self.assertIn('results/execution.json',z.namelist())
            self.assertIn('paper/main.md',z.namelist())
            self.assertNotIn('researchdesk.sqlite3',z.namelist())
        self.assertFalse(any(r['step'].startswith('api:') for r in self.store.history(pid)[1]))

    def test_all_three_protocols_and_usage(self):
        for protocol in ('chat','responses','anthropic'):
            settings=providers.validate_settings({'protocol':protocol,'endpoint':self.base,'model':'test-model','role_models':{'reviewer':'review-model'},'input_price':2,'output_price':4})
            text,usage=providers.call(settings,'local-fake-key','system','hello','reviewer')
            self.assertEqual(text,'test response')
            self.assertEqual(usage['input_tokens'],100)
            self.assertEqual(usage['output_tokens'],50)
            self.assertAlmostEqual(usage['estimated_cost'],.0004)
            request=FakeAPI.records[-1]
            self.assertEqual(request['body']['model'],'review-model')
            if protocol=='responses':
                self.assertIn('max_output_tokens',request['body']);self.assertFalse(request['body']['store'])
            if protocol=='anthropic': self.assertEqual(request['headers']['X-Api-Key'],'local-fake-key')

    def test_fake_api_pipeline(self):
        self.store.save_settings({'endpoint':self.base,'model':'fake-model'})
        pid=self.project('api')
        self.run_to_end(pid)
        receipts=self.store.history(pid)[1]
        self.assertEqual(len([r for r in receipts if r['step'].startswith('api:')]),10)
        self.assertTrue(any(f['path'].startswith('audit/') for f in self.store.files(pid)))

    def test_writing_template_no_computation(self):
        pid=self.project(template='writing')
        self.engine.start(pid)
        self.assertEqual(self.wait(pid)['status'],'review')
        self.engine.start(pid)
        self.assertEqual(self.wait(pid)['status'],'completed')
        self.assertFalse(self.store.path(pid,'results/metrics.json').exists())
        self.assertIn('没有实验数值',self.store.path(pid,'paper/main.md').read_text('utf-8'))

    def test_explicit_consent_required(self):
        pid=self.project()
        self.engine.start(pid);self.wait(pid)
        self.engine.start(pid);self.wait(pid)
        with self.assertRaisesRegex(ValueError,'审阅'): self.engine.start(pid)
        self.assertEqual(self.store.get(pid)['status'],'awaiting_execution')

    def test_failed_execution_and_skip_never_reuses_metrics(self):
        pid=self.project()
        self.engine.start(pid);self.wait(pid)
        self.engine.start(pid);self.wait(pid)
        bad="from pathlib import Path\nPath('results/metrics.json').write_text('{\"metrics\": []}')\nraise RuntimeError('failed')\n"
        self.store.write(pid,'code/experiment.py',bad)
        self.engine.start(pid,'execute')
        self.assertEqual(self.wait(pid)['status'],'awaiting_execution')
        self.assertFalse(self.store.path(pid,'results/metrics.json').exists())
        self.assertTrue(self.store.path(pid,'results/execution.log').exists())
        self.engine.start(pid,'skip')
        self.assertEqual(self.wait(pid)['status'],'review')
        self.assertFalse(self.store.path(pid,'results/metrics.png').exists())
        self.assertIn('没有实验数值',self.store.path(pid,'paper/main.md').read_text('utf-8'))

    def test_invalid_model_outputs_fail_and_can_retry(self):
        pid=self.project('api')
        self.store.save_settings({'endpoint':self.base,'model':'fake'})
        with patch('researchdesk.providers.call',return_value=('not json',{'model':'fake'})):
            self.engine.start(pid)
            self.assertEqual(self.wait(pid)['status'],'failed')
        self.assertEqual(self.store.get(pid)['cursor'],0)
        self.assertFalse(self.store.path(pid,'docs/brief.md').exists())
        self.engine.start(pid)
        self.assertEqual(self.wait(pid)['status'],'review')

    def test_pause_and_duplicate_task(self):
        pid=self.project()
        entered=threading.Event();released=threading.Event()
        original=self.engine.generate
        def slow(*args):
            entered.set();released.wait(3);return original(*args)
        with patch.object(self.engine,'generate',side_effect=slow):
            self.engine.start(pid)
            self.assertTrue(entered.wait(3))
            with self.assertRaisesRegex(ValueError,'已有任务'): self.engine.start(pid)
            self.engine.pause(pid);released.set()
            self.assertEqual(self.wait(pid)['status'],'paused')
        self.assertEqual(self.store.get(pid)['cursor'],1)
        self.engine.start(pid)
        self.assertEqual(self.wait(pid)['status'],'review')

    def test_restart_recovers_code_confirmation(self):
        pid=self.project()
        self.store.update(pid,status='running',cursor=5)
        recovered=Store(self.tmp.name)
        self.assertEqual(recovered.get(pid)['status'],'awaiting_execution')

    def test_path_traversal_and_generated_contract(self):
        pid=self.project()
        for path in ('../outside','/Windows/test','C:/x','docs/../../settings','docs/evil\\file','.key'):
            with self.subTest(path=path),self.assertRaises(ValueError): self.store.path(pid,path)
        self.store.save_settings({'endpoint':self.base,'model':'fake'})
        pid=self.project('api')
        raw=json.dumps({'files':[{'path':'../evil.txt','content':'evil'}]})
        with patch('researchdesk.providers.call',return_value=(raw,{})),self.assertRaisesRegex(ValueError,'产物路径'):
            self.engine.generate(pid,BY_ID['brief'])

    def test_key_encrypted_and_not_returned(self):
        secret='sentinel-test-secret-8127634'
        self.store.save_settings({'endpoint':self.base,'model':'fake','api_key':secret})
        self.assertNotIn(secret,json.dumps(self.store.settings()))
        self.assertNotIn(secret.encode(),self.store.db.read_bytes())
        self.assertTrue(self.store.settings()['has_key'])
        self.assertEqual(self.store.settings(private=True)['key'],secret)
        if self.store.settings()['key_storage']=='session':
            self.assertFalse(Store(self.tmp.name).settings()['has_key'])
        self.store.save_settings({'endpoint':self.base,'model':'fake','clear_key':True})
        self.assertFalse(self.store.settings()['has_key'])

    def test_invalid_metrics_are_rejected(self):
        for value in (True,None,'0.4',math.nan,math.inf):
            with self.subTest(value=value),self.assertRaises(ValueError): validate_metrics({'metrics':[{'name':'m','value':value}],'dataset_origin':'simulation','method':'test'})

class HTTPTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        self.server.app=Application(self.tmp.name)
        self.base=f'http://127.0.0.1:{self.server.server_port}'
        threading.Thread(target=self.server.serve_forever,daemon=True).start()
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.tmp.cleanup()
    def request(self,path,data=None,headers=None):
        headers=headers or {}
        req=urllib.request.Request(self.base+path,json.dumps(data).encode() if data is not None else None,headers)
        return urllib.request.urlopen(req,timeout=10)
    def test_csrf_host_and_settings_redaction(self):
        with self.request('/api/bootstrap') as r: boot=json.load(r)
        with self.assertRaises(urllib.error.HTTPError) as cm: self.request('/api/projects',{})
        self.assertEqual(cm.exception.code,403)
        with self.assertRaises(urllib.error.HTTPError) as cm: self.request('/api/bootstrap',headers={'Host':'evil.example'})
        self.assertEqual(cm.exception.code,403)
        headers={'X-ResearchDesk-Token':boot['token'],'Content-Type':'application/json'}
        with self.request('/api/projects',{'title':'HTTP测试','goal':'测试','mode':'api'},headers) as r: p=json.load(r)
        self.assertEqual(p['status'],'draft')
        with self.assertRaises(urllib.error.HTTPError) as cm: self.request('/api/projects',{},dict(headers,Origin='https://evil.example'))
        self.assertEqual(cm.exception.code,403)
        with self.request('/') as r:
            self.assertIn('ResearchDesk',r.read().decode());self.assertIn("object-src 'none'",r.headers['Content-Security-Policy'])

    def test_upload_text_word_pdf_and_freeze(self):
        app=self.server.app
        pid=app.store.create('资料测试','research','分析资料','api')['id']
        text='col,value\n测试,1\n'
        app.upload(pid,{'name':'数据.csv','content':base64.b64encode(text.encode()).decode()})
        self.assertEqual(app.store.path(pid,'materials/数据.csv').read_text('utf-8'),text)
        from docx import Document
        import io
        buf=io.BytesIO();doc=Document();doc.add_paragraph('真实导入文本');doc.save(buf)
        app.upload(pid,{'name':'资料.docx','content':base64.b64encode(buf.getvalue()).decode()})
        self.assertIn('真实导入文本',app.store.path(pid,'materials/资料.docx.extracted.txt').read_text('utf-8'))
        from reportlab.pdfgen import canvas
        buf=io.BytesIO();pdf=canvas.Canvas(buf);pdf.drawString(50,700,'Research test evidence');pdf.save()
        app.upload(pid,{'name':'paper.pdf','content':base64.b64encode(buf.getvalue()).decode()})
        self.assertIn('Research test evidence',app.store.path(pid,'materials/paper.pdf.extracted.txt').read_text('utf-8'))
        app.store.update(pid,status='review')
        with self.assertRaisesRegex(ValueError,'启动流程前'): app.upload(pid,{'name':'new.txt','content':'eA=='})

if __name__=='__main__': unittest.main(verbosity=2)
